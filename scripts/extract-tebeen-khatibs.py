"""Prepare Tebeen khatibs, preferences and explicit permanent assignments from the workbook.

Output contains private phone numbers: keep it outside version control (normally scratch/).
"""
import argparse
import csv
import json
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


def text(value):
    if value is None:
        return ''
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def norm(value):
    return re.sub(r'\s+', ' ', text(value).lower().translate(str.maketrans('أإآىةؤئ', 'ااايهوي'))).strip()


def phone(value):
    raw = text(value)
    if not raw:
        return None
    digits = re.sub(r'[\s\-()+\u200e\u200f\u202a-\u202e\u2066-\u2069]', '', raw.translate(str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹', '01234567890123456789')))
    if digits.startswith('0020'):
        digits = '0' + digits[4:]
    elif digits.startswith('20') and len(digits) == 12:
        digits = '0' + digits[2:]
    elif len(digits) == 10 and digits.startswith('1'):
        digits = '0' + digits
    if not re.fullmatch(r'01\d{9}', digits):
        raise ValueError('Phone number needs manual review')
    return digits


def quote(value):
    if value is None:
        return 'NULL'
    # Keep note line breaks exact even if the SQL importer normalizes Windows lines.
    return ' || char(10) || '.join("'" + part.replace("'", "''") + "'" for part in str(value).split('\n'))


def extract(workbook_path, rest_type, root):
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=False)
    with (root / 'context/tebeen-mosques.csv').open(encoding='utf-8', newline='') as source:
        mosques = list(csv.DictReader(source))
    label_map = {norm(m['name'] + ' ' + m['mujawra']): m['id'] for m in mosques}
    assert len(label_map) == len(mosques)
    sheet = workbook['الخطباء']
    assert norm(sheet.cell(1, 1).value) == norm('الإسم')
    assert norm(sheet.cell(1, 2).value) == norm('رقم الهاتف')
    records, by_name, issues = [], {}, []
    for cells in sheet.iter_rows(min_row=2):
        if not isinstance(cells[0].value, str) or not text(cells[0].value) or text(cells[0].value) == '0':
            continue
        name = text(cells[0].value)
        key = norm(name)
        if key in by_name:
            raise ValueError(f'Duplicate normalized khatib identity at row {cells[0].row}')
        notes = text(cells[2].value)
        inactive = any(norm(label) in norm(notes) for label in ['متوقف مؤقتا', 'ممنوع حاليا'])
        preferences = []
        for cell in cells[43:]:
            if not isinstance(cell.value, str) or not text(cell.value) or text(cell.value) == '0':
                continue
            label = norm(cell.value)
            if label not in label_map:
                raise ValueError(f'Unmatched preference at {cell.coordinate}')
            if label_map[label] not in preferences:
                preferences.append(label_map[label])
        record = {'id': f'K{len(records)+1:03d}', 'name': name, 'phone': phone(cells[1].value),
                  'type': 'primary' if norm(notes) == norm('امام') else rest_type,
                  'notes': notes, 'active': int(not inactive), 'preferences': preferences,
                  'sourceRow': cells[0].row}
        records.append(record)
        by_name[key] = record
    duplicates = defaultdict(list)
    for record in records:
        if record['phone']:
            duplicates[record['phone']].append(record)
    source_phone_count = sum(bool(record['phone']) for record in records)
    for number, group in duplicates.items():
        if len(group) < 2:
            continue
        issues.append({'kind': 'shared_phone', 'khatibs': [r['name'] for r in group]})
        for record in group:
            record['notes'] = '\n'.join(filter(None, [record['notes'], 'رقم الهاتف الوارد بالشيت مكرر ويحتاج مراجعة: ' + number]))
            record['phone'] = None
    source = workbook['المساجد']
    permanent_column = [cell.column for cell in source[1] if norm(cell.value) == norm('خطيب ثابت')]
    assert len(permanent_column) == 1
    candidates = []
    for mosque in mosques:
        row = int(mosque['sort_order']) + 1
        assert norm(source.cell(row, 2).value) == norm(mosque['name'])
        assert norm(source.cell(row, 3).value) == norm(mosque['mujawra'])
        name = text(source.cell(row, permanent_column[0]).value)
        if not name:
            continue
        if norm(name) not in by_name:
            raise ValueError(f'Unknown permanent khatib at row {row}')
        candidates.append((mosque, by_name[norm(name)]))
    frequency = Counter(k['id'] for m, k in candidates)
    permanent = []
    for mosque, record in candidates:
        reason = None
        if frequency[record['id']] > 1:
            reason = 'same_khatib_in_multiple_mosques'
        elif not record['active']:
            reason = 'inactive_khatib'
        elif norm('لايخطب في') in norm(record['notes']) and norm(mosque['name']) in norm(record['notes']):
            reason = 'explicit_note_excludes_mosque'
        if reason:
            issues.append({'kind': reason, 'mosqueId': mosque['id'], 'mosque': mosque['name'],
                           'area': mosque['mujawra'], 'khatib': record['name']})
        else:
            permanent.append({'mosqueId': mosque['id'], 'khatibId': record['id']})
    workbook.close()
    return {'khatibs': records, 'permanent': permanent, 'issues': issues, 'summary': {
        'khatibs': len(records), 'sourcePhones': source_phone_count,
        'phoneFields': sum(bool(r['phone']) for r in records),
        'preferences': sum(len(r['preferences']) for r in records),
        'khatibsWithPreferences': sum(bool(r['preferences']) for r in records),
        'inactive': sum(not r['active'] for r in records),
        'permanent': len(permanent), 'permanentCandidates': len(candidates),
        'types': dict(Counter(r['type'] for r in records))}}


def prepare_sql(data):
    now = datetime.now(timezone.utc).isoformat(timespec='milliseconds').replace('+00:00', 'Z')
    statements = ['-- First import into khotba-tebeen-db ONLY. Contains private phone numbers.']
    for record in data['khatibs']:
        values = [record['id'], record['name'], record['phone'], record['type'], record['notes'], record['active'], now]
        statements.append('INSERT INTO khatibs(id,name,phone,type,notes,active,created_at) VALUES (' + ','.join(map(quote, values)) + ');')
        for mid in record['preferences']:
            statements.append(f"INSERT INTO preferences(khatib_id,mosque_id) VALUES ({quote(record['id'])},{quote(mid)});")
    for pair in data['permanent']:
        statements.append(f"UPDATE mosques SET permanent_khatib_id={quote(pair['khatibId'])} WHERE id={quote(pair['mosqueId'])} AND permanent_khatib_id IS NULL;")
    statements.append(f"INSERT INTO settings(key,value) VALUES ('next_khatib_seq','{len(data['khatibs'])+1}') ON CONFLICT(key) DO UPDATE SET value=excluded.value;")
    return '\n'.join(statements) + '\n'


def verify_local(root, data, sql):
    db = sqlite3.connect(':memory:')
    db.executescript((root / 'worker/schema.sql').read_text(encoding='utf-8'))
    with (root / 'context/tebeen-mosques.csv').open(encoding='utf-8', newline='') as source:
        for m in csv.DictReader(source):
            db.execute('INSERT INTO mosques(id,name,mujawra,sort_order,created_at) VALUES (?,?,?,?,?)',
                       (m['id'], m['name'], m['mujawra'], m['sort_order'], '2026-10-20'))
    db.commit()
    db.executescript(sql)
    assert db.execute('SELECT COUNT(*) FROM khatibs').fetchone()[0] == len(data['khatibs'])
    assert db.execute('SELECT COUNT(*) FROM preferences').fetchone()[0] == data['summary']['preferences']
    assert db.execute('SELECT COUNT(*) FROM mosques WHERE permanent_khatib_id IS NOT NULL').fetchone()[0] == len(data['permanent'])
    assert db.execute('SELECT COUNT(*) FROM assignments').fetchone()[0] == 0
    assert not db.execute('PRAGMA foreign_key_check').fetchall()
    for r in data['khatibs']:
        actual = db.execute('SELECT name,phone,type,notes,active FROM khatibs WHERE id=?', (r['id'],)).fetchone()
        assert actual == tuple(r[k] for k in ['name', 'phone', 'type', 'notes', 'active'])
    db.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--rest-type', choices=['primary', 'volunteer'], required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parent.parent
    data = extract(args.workbook, args.rest_type, root)
    sql = prepare_sql(data)
    verify_local(root, data, sql)
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output / 'tebeen-khatibs.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    (args.output / 'tebeen-khatibs.sql').write_text(sql, encoding='utf-8')
    print(json.dumps(data['summary'], ensure_ascii=False))
    print(json.dumps(data['issues'], ensure_ascii=False))


if __name__ == '__main__':
    main()
